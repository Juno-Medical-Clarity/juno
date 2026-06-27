# Dataset: MedicationQA — Gold standard consumer medication QA corpus
# Source: https://github.com/abachaa/Medication_QA_MedInfo2019
# Paper: "Bridging the Gap between Consumers' Medication Questions and Trusted Answers" (MedInfo 2019)
# License: CC BY 4.0
#
# Excel file: MedInfo2019-QA-Medications.xlsx (sheet: DrugQA)
# Columns: Question, Focus (Drug), Question Type, Answer, Section Title, URL
# 690 question-answer pairs about consumer medication questions.

import os
import sys
import urllib.request
import openpyxl

# ── Config ────────────────────────────────────────────────────────────────────

DATASET_URL = (
    "https://raw.githubusercontent.com/abachaa/Medication_QA_MedInfo2019"
    "/master/MedInfo2019-QA-Medications.xlsx"
)
OUTPUT_DIR = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "preset-data", "medicationqa"
)
TMP_XLSX = "/tmp/MedInfo2019-QA-Medications.xlsx"

# ── Download ──────────────────────────────────────────────────────────────────

def download_xlsx(url: str, dest: str) -> None:
    if os.path.exists(dest):
        print(f"Using cached file: {dest}")
        return
    print(f"Downloading dataset from {url} …")
    urllib.request.urlretrieve(url, dest)
    print(f"Saved to {dest}")

# ── Parse & save ──────────────────────────────────────────────────────────────

def save_text(path: str, content: str) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content.strip() + "\n")

def parse_and_save(xlsx_path: str, out_dir: str) -> int:
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["DrugQA"]

    # Validate headers
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    expected = ["Question", "Focus (Drug)", "Question Type", "Answer", "Section Title", "URL"]
    if headers != expected:
        print(f"WARNING: unexpected headers: {headers}", file=sys.stderr)

    count = 0
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        question, drug_focus, question_type, answer, section_title, source_url = row

        # Skip rows with missing question or answer
        if not question or not answer:
            print(f"  Skipping row {idx}: missing question or answer")
            continue

        sample_id = f"{idx:04d}"
        sample_dir = os.path.join(out_dir, sample_id)

        save_text(os.path.join(sample_dir, "question.txt"), str(question))
        save_text(os.path.join(sample_dir, "answer.txt"), str(answer))

        # Build metadata
        meta_lines = []
        if drug_focus:
            meta_lines.append(f"drug_focus: {drug_focus}")
        if question_type:
            meta_lines.append(f"question_type: {question_type}")
        if section_title:
            meta_lines.append(f"section_title: {section_title}")
        if source_url:
            meta_lines.append(f"source_url: {source_url}")
        meta_lines.append(f"sample_id: {sample_id}")

        if meta_lines:
            save_text(os.path.join(sample_dir, "metadata.txt"), "\n".join(meta_lines))

        count += 1

    return count

# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    out_dir = os.path.realpath(OUTPUT_DIR)
    os.makedirs(out_dir, exist_ok=True)

    download_xlsx(DATASET_URL, TMP_XLSX)
    count = parse_and_save(TMP_XLSX, out_dir)

    print(f"\nDone. Saved {count} samples to {out_dir}")
    print(f"Example layout:")
    print(f"  {out_dir}/0001/question.txt")
    print(f"  {out_dir}/0001/answer.txt")
    print(f"  {out_dir}/0001/metadata.txt")

if __name__ == "__main__":
    main()
